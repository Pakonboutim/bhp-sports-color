import os, json, re, uuid
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
import requests

XLSX = os.environ.get("SPORTS_XLSX", r"C:\Users\pakon\Downloads\กีฬาสีโรงเรียนบ้านห้วยผึ้งปี 2569.xlsx")
URL = os.environ["SUPABASE_URL"].rstrip("/")
KEY = os.environ["SUPABASE_SECRET_KEY"]

BASE_HEADERS = {
    "apikey": KEY,
    "Authorization": f"Bearer {KEY}",
    "Content-Type": "application/json",
}

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v

def code(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and float(v).is_integer():
        return str(int(v))
    return str(v).strip()

def number(v):
    if v is None or v == "":
        return None
    return int(float(v))

def room(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)) and float(v).is_integer():
        return str(int(v))
    return str(v).strip()

def iso_dt(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    return str(v).strip()

def iso_date(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        y = int(m.group(1))
        if y > 2400:
            y -= 543
        return f"{y:04d}-{m.group(2)}-{m.group(3)}"
    return s

def iso_time(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.strftime("%H:%M:%S")
    s = str(v).strip()
    if re.fullmatch(r"\d{1,2}:\d{2}", s):
        h, m = s.split(":")
        return f"{int(h):02d}:{m}:00"
    return s

def rows(ws):
    vals = list(ws.iter_rows(values_only=True))
    if not vals:
        return []
    headers = [str(x).strip() if x is not None else "" for x in vals[0]]
    out = []
    for row in vals[1:]:
        if not any(v is not None and str(v).strip() != "" for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return out

def request(method, path, *, params=None, body=None, prefer=None):
    headers = dict(BASE_HEADERS)
    if prefer:
        headers["Prefer"] = prefer
    r = requests.request(method, f"{URL}/rest/v1/{path}", headers=headers, params=params, json=body, timeout=120)
    if not r.ok:
        raise RuntimeError(f"{method} {path} failed: {r.status_code} {r.text}")
    return r

def insert_rows(table, data, batch_size=200):
    if not data:
        print(f"{table}: 0 rows (skip)")
        return
    for i in range(0, len(data), batch_size):
        batch = data[i:i+batch_size]
        request("POST", table, body=batch, prefer="return=minimal")
        print(f"{table}: {min(i+batch_size, len(data))}/{len(data)}")

def upsert(table, data, conflict, batch_size=200):
    if not data:
        print(f"{table}: 0 rows (skip)")
        return
    for i in range(0, len(data), batch_size):
        batch = data[i:i+batch_size]
        request("POST", table, params={"on_conflict": conflict}, body=batch, prefer="resolution=merge-duplicates,return=minimal")
        print(f"{table}: {min(i+batch_size, len(data))}/{len(data)}")

def count_rows(table):
    headers = dict(BASE_HEADERS)
    headers["Prefer"] = "count=exact"
    r = requests.get(f"{URL}/rest/v1/{table}", headers=headers, params={"select":"*"}, timeout=60)
    r.raise_for_status()
    return int(r.headers.get("content-range","*/0").split("/")[-1])

def get_students_lookup():
    r = request("GET", "students", params={"select":"student_id,student_code,first_name,last_name,level,room,number"})
    rows_ = r.json()
    by_code = {}
    for s in rows_:
        by_code.setdefault(str(s["student_code"]), []).append(s)
    return by_code

def resolve_student(by_code, student_code, student_name=None, level_room=None):
    matches = by_code.get(str(student_code), [])
    if len(matches) == 1:
        return matches[0]["student_id"]
    if not matches:
        raise RuntimeError(f"ไม่พบนักเรียนรหัส {student_code}")

    # Legacy athlete/photo records may use duplicated student codes.
    # Try to disambiguate by name and level/room if supplied.
    target_name = (student_name or "").replace(" ","")
    target_lr = (level_room or "").replace(" ","")
    filtered = matches
    if target_name:
        filtered = [s for s in filtered if f"{s.get('first_name','')}{s.get('last_name','')}".replace(" ","") in target_name or target_name in f"{s.get('first_name','')}{s.get('last_name','')}".replace(" ","")]
    if target_lr and len(filtered) != 1:
        filtered2 = [s for s in filtered if f"{s.get('level','')}/{s.get('room','')}".replace(" ","") == target_lr]
        if filtered2:
            filtered = filtered2
    if len(filtered) == 1:
        return filtered[0]["student_id"]
    raise RuntimeError(f"รหัสนักเรียน {student_code} ซ้ำและระบุคนไม่ได้: {matches}")

p = Path(XLSX)
if not p.exists():
    raise SystemExit(f"ไม่พบไฟล์ Excel: {p}")

# Protect against accidental duplicate imports.
existing = count_rows("students")
if existing != 0:
    raise SystemExit(f"หยุด: ตาราง students มีข้อมูลอยู่แล้ว {existing} แถว กรุณาตรวจสอบก่อนรันซ้ำ")

wb = load_workbook(p, data_only=True, read_only=True)

students = []
for r in rows(wb["นักเรียน"]):
    students.append({
        "student_id": str(uuid.uuid4()),
        "student_code": code(r.get("เลขประจำตัว")),
        "number": number(r.get("เลขที่")),
        "prefix": clean(r.get("คำนำหน้า")),
        "first_name": clean(r.get("ชื่อ")) or "",
        "last_name": clean(r.get("สกุล")) or "",
        "level": clean(r.get("ชั้น")) or "",
        "room": room(r.get("ห้อง")),
        "color": clean(r.get("สี")),
        "active": True,
    })

teachers = []
for r in rows(wb["ครูกีฬาสี"]):
    teachers.append({
        "color": clean(r.get("สี")),
        "prefix": clean(r.get("คำนำหน้า")),
        "first_name": clean(r.get("ชื่อ")) or "",
        "last_name": clean(r.get("นามสกุล")) or "",
        "role": clean(r.get("ตำแหน่ง/หมายเหตุ")),
        "active": True,
    })

sports = []
for r in rows(wb["Sports"]):
    sports.append({
        "sport_id": clean(r.get("SportId")),
        "name": clean(r.get("Name")) or "",
        "level": clean(r.get("Level")),
        "gender": clean(r.get("Gender")) or "Mixed",
        "athlete_limit": number(r.get("AthleteLimit")) or 1,
        "type": clean(r.get("Type")) or "Knockout",
        "active": bool(r.get("Active")) if r.get("Active") is not None else True,
        "team_format": clean(r.get("TeamFormat")) or "Standard4",
        "updated_at": iso_dt(r.get("UpdatedAt")),
    })

matches = []
for r in rows(wb["Matches"]):
    status = clean(r.get("Status")) or "Pending"
    score_a = number(r.get("ScoreA"))
    score_b = number(r.get("ScoreB"))
    ts = iso_dt(r.get("Timestamp"))
    matches.append({
        "match_id": clean(r.get("MatchId")),
        "sport_id": clean(r.get("SportId")),
        "sport_name": clean(r.get("SportName")) or "",
        "round": clean(r.get("Round")) or "",
        "team_a": clean(r.get("TeamA")),
        "team_b": clean(r.get("TeamB")),
        "score_a": score_a,
        "score_b": score_b,
        "winner": clean(r.get("Winner")),
        "loser": clean(r.get("Loser")),
        "referee_name": clean(r.get("Referee")),
        "confirmed_at": ts if status == "Confirmed" else None,
        "status": status,
        "match_date": iso_date(r.get("MatchDate")),
        "match_time": iso_time(r.get("MatchTime")),
    })

settings = []
for r in rows(wb["Settings"]):
    key = clean(r.get("Key"))
    if not key:
        continue
    val = clean(r.get("Value"))
    if isinstance(val, bool):
        json_val = val
    elif isinstance(val, (int, float)):
        json_val = val
    else:
        s = "" if val is None else str(val).strip()
        json_val = (s.lower() == "true") if s.lower() in ("true","false") else s
    settings.append({
        "key": key,
        "value": json_val,
        "updated_at": iso_dt(r.get("UpdatedAt")) or datetime.now().isoformat(),
    })

audit = []
for r in rows(wb["AuditLog"]):
    details = clean(r.get("Details"))
    if isinstance(details, str):
        try:
            details = json.loads(details)
        except Exception:
            details = {"raw": details}
    elif details is None:
        details = {}
    audit.append({
        "event_time": iso_dt(r.get("Time")) or datetime.now().isoformat(),
        "action": clean(r.get("Action")) or "UNKNOWN",
        "user_type": clean(r.get("UserType")) or "System",
        "details": details,
    })

print("ข้อมูลที่พบใน Excel")
print(f"  students: {len(students)}")
print(f"  teachers: {len(teachers)}")
print(f"  sports: {len(sports)}")
print(f"  matches: {len(matches)}")
print(f"  audit_logs: {len(audit)}")

# Students must be inserted, not upserted by student_code, because codes can legitimately duplicate.
insert_rows("students", students)

# Build lookup after inserting students.
by_code = get_students_lookup()

athletes = []
for r in rows(wb["Athletes"]):
    sid = code(r.get("StudentId"))
    if not sid:
        continue
    student_id = resolve_student(by_code, sid, clean(r.get("StudentName")), clean(r.get("LevelRoom")))
    athletes.append({
        "sport_id": clean(r.get("SportId")),
        "color": clean(r.get("Color")),
        "student_id": student_id,
        "student_code": sid,
        "student_name": clean(r.get("StudentName")) or "",
        "level_room": clean(r.get("LevelRoom")) or "",
        "updated_at": iso_dt(r.get("UpdatedAt")),
    })

photos = []
for r in rows(wb["StudentPhotos"]):
    sid = code(r.get("StudentId"))
    if not sid:
        continue
    student_id = resolve_student(by_code, sid)
    photos.append({
        "student_id": student_id,
        "student_code": sid,
        "legacy_drive_file_id": clean(r.get("DriveFileId")),
        "legacy_photo_url": clean(r.get("PhotoUrl")),
        "photo_url": clean(r.get("PhotoUrl")),
        "updated_at": iso_dt(r.get("UpdatedAt")) or datetime.now().isoformat(),
    })

upsert("teachers", teachers, "color,prefix,first_name,last_name")
upsert("sports", sports, "sport_id")
upsert("athletes", athletes, "sport_id,color,student_id")
upsert("matches", matches, "match_id")
upsert("student_photos", photos, "student_id")
upsert("app_settings", settings, "key")

if audit:
    insert_rows("audit_logs", audit)

print("\nVERIFY")
for table in ["students","teachers","sports","athletes","matches","student_photos","app_settings","audit_logs"]:
    print(f"{table}: {count_rows(table)}")

print("\nIMPORT COMPLETE")
